"""
migrar_planilha.py
─────────────────────────────────────────────────────────────────────
Migra os dados históricos da planilha Excel (aba Base + Cadastro) para o
banco Postgres/SQLite do novo dashboard.

Fonte de verdade dos nomes de unidade e do histórico de faturas: a aba
`Base` (coluna "Ref. Local"), a mesma usada por toda a automação em
Python já existente (Agente Equatorial/orquestrador_mensal.py). Os
códigos de UC de cada unidade são derivados de TODOS os valores já
observados na própria aba Base (uma unidade pode ter mais de um código
ao longo do tempo, por causa da troca de formato ANEEL em 01/04/2026).

Uso:
    python3 migrar_planilha.py /caminho/para/VClaude V2.xlsx

Variável de ambiente DATABASE_URL define o banco de destino (default:
sqlite local ./dev.db, útil para testar antes de apontar pro Postgres
do Railway).
─────────────────────────────────────────────────────────────────────
"""

import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.db import Base, SessionLocal, engine
from app.models import Fatura, Investimento, Unidade, UnidadeCodigo

# Nomes de unidades geradoras (não entram no rateio de consumo)
GERADORAS = {"Sonho Verde"}

_avisos_conversao = []


def to_float_seguro(valor, contexto=""):
    """
    Alguns lançamentos manuais na planilha guardaram número como texto
    malformado (ex: '1,821,81' em vez de 1821.81 — vírgula usada tanto como
    separador de milhar quanto decimal). Converte com robustez e registra um
    aviso em vez de derrubar a migração.
    """
    if valor is None or valor == "":
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    if isinstance(valor, str):
        s = valor.strip()
        if "," in s:
            partes = s.split(",")
            s = "".join(partes[:-1]) + "." + partes[-1]
        try:
            resultado = float(s)
            _avisos_conversao.append(
                f"{contexto}: valor texto '{valor}' convertido para {resultado}"
            )
            return resultado
        except ValueError:
            _avisos_conversao.append(
                f"{contexto}: valor texto '{valor}' NÃO PÔDE ser convertido, usando 0.0"
            )
            return 0.0
    return 0.0


def carregar_cadastro(ws_cad):
    """Retorna dict nome -> {tipo, status, pct_atual, pct_sugerido}."""
    dados = {}
    for r in range(4, ws_cad.max_row + 1):
        nome = ws_cad.cell(r, 1).value
        if not nome or str(nome).strip().upper() == "TOTAL":
            continue
        dados[str(nome).strip()] = {
            "tipo": ws_cad.cell(r, 5).value,
            "status": ws_cad.cell(r, 3).value or "Ativa",
            "pct_atual": ws_cad.cell(r, 7).value or 0.0,
            "pct_sugerido": ws_cad.cell(r, 8).value or 0.0,
        }
    return dados


def carregar_investimento(ws_painel):
    """Parcelas de investimento (Painel!A5:B8)."""
    itens = []
    for r in range(5, 9):
        desc = ws_painel.cell(r, 1).value
        valor = ws_painel.cell(r, 2).value
        if desc and isinstance(valor, (int, float)):
            itens.append((str(desc).strip(), float(valor)))
    return itens


def migrar(planilha_path: str):
    print(f"Lendo planilha: {planilha_path}")
    wb = load_workbook(planilha_path, data_only=True)
    ws_base = wb["Base"]
    ws_cad = wb["Cadastro"] if "Cadastro" in wb.sheetnames else None
    ws_painel = wb["Painel"] if "Painel" in wb.sheetnames else None

    cadastro = carregar_cadastro(ws_cad) if ws_cad else {}

    # 1ª passada: descobrir todas as unidades (Ref. Local) e códigos observados
    codigos_por_unidade = defaultdict(set)
    linhas_por_unidade = defaultdict(list)

    for row in ws_base.iter_rows(min_row=2, values_only=True):
        codigo, local = row[0], row[1]
        ano_mes = row[3]
        if not local or not isinstance(ano_mes, datetime):
            continue
        local = str(local).strip()
        if codigo:
            codigos_por_unidade[local].add(str(codigo).strip())
        linhas_por_unidade[local].append(row)

    print(f"Unidades encontradas na aba Base: {len(linhas_por_unidade)}")

    Base.metadata.create_all(bind=engine)
    db = SessionLocal()

    try:
        unidade_obj = {}
        for nome in linhas_por_unidade:
            info = cadastro.get(nome, {})
            u = Unidade(
                nome=nome,
                tipo=info.get("tipo"),
                status=info.get("status") or "Ativa",
                geradora=(nome in GERADORAS),
                pct_rateio_atual=float(info.get("pct_atual") or 0.0),
                pct_rateio_sugerido=float(info.get("pct_sugerido") or 0.0),
            )
            db.add(u)
            unidade_obj[nome] = u
        db.flush()  # garante IDs antes de criar filhos

        total_codigos = 0
        for nome, codigos in codigos_por_unidade.items():
            for codigo in codigos:
                db.add(UnidadeCodigo(unidade_id=unidade_obj[nome].id, codigo=codigo))
                total_codigos += 1
        print(f"Códigos de UC migrados: {total_codigos}")

        total_faturas = 0
        ignoradas_zero = 0
        for nome, linhas in linhas_por_unidade.items():
            u = unidade_obj[nome]
            for row in linhas:
                (codigo, local, ano, ano_mes, kwh_tax, cip_tax, consumo,
                 injetada, difkw, acumulado, expirar, valor, pago, *_resto) = (
                    list(row) + [None] * (14 - len(row))
                )
                # pula linhas futuras/placeholder totalmente zeradas
                if not consumo and not valor and not pago:
                    ignoradas_zero += 1
                    continue
                ctx = f"{nome} {ano_mes:%m/%Y}"
                consumo_f   = to_float_seguro(consumo, f"{ctx} (REAL CONSUMO)")
                injetada_f  = to_float_seguro(injetada, f"{ctx} (Energia Injetada)")
                difkw_f     = to_float_seguro(difkw, f"{ctx} (Diferença KW)")
                acumulado_f = to_float_seguro(acumulado, f"{ctx} (Acumulado)")
                expirar_f   = to_float_seguro(expirar, f"{ctx} (A expirar)")
                valor_f     = to_float_seguro(valor, f"{ctx} (REAL VALOR)")
                pago_f      = to_float_seguro(pago, f"{ctx} (PAGO)")
                diferenca   = valor_f - pago_f
                db.add(Fatura(
                    unidade_id=u.id,
                    ano_mes=ano_mes.date(),
                    kwh_tax=to_float_seguro(kwh_tax, f"{ctx} (KWh tax)"),
                    cip_tax=to_float_seguro(cip_tax, f"{ctx} (Cip tax)"),
                    real_consumo=consumo_f,
                    energia_injetada=injetada_f,
                    diferenca_kw=difkw_f,
                    acumulado=acumulado_f,
                    a_expirar=expirar_f,
                    real_valor=valor_f,
                    pago=pago_f,
                    diferenca=diferenca,
                ))
                total_faturas += 1

        print(f"Faturas migradas: {total_faturas} (ignoradas por estarem vazias: {ignoradas_zero})")

        if ws_painel:
            itens = carregar_investimento(ws_painel)
            for desc, valor in itens:
                db.add(Investimento(descricao=desc, valor=valor))
            print(f"Parcelas de investimento migradas: {len(itens)}")

        db.commit()
        print("✅ Migração concluída.")

        if _avisos_conversao:
            print(f"\n⚠️  {len(_avisos_conversao)} valor(es) com formato de texto malformado "
                  f"na planilha (corrigidos automaticamente na migração, mas vale revisar "
                  f"a fonte na aba Base):")
            for aviso in _avisos_conversao:
                print(f"   - {aviso}")
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Uso: python3 migrar_planilha.py /caminho/para/planilha.xlsx")
        sys.exit(1)
    caminho = Path(sys.argv[1])
    if not caminho.exists():
        print(f"❌ Arquivo não encontrado: {caminho}")
        sys.exit(1)
    migrar(str(caminho))
